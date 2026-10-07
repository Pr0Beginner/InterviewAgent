"""解析上传的简历，并将纯文本安全地保存在本地。"""

from __future__ import annotations

import os
import re
import unicodedata
from datetime import datetime
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import numpy as np
import pymupdf
from docx import Document
from pypdf import PdfReader

from backend.app.core.config import get_settings
from backend.app.core.exceptions import ApplicationError
from backend.app.storage.local import LocalStateStore


class ResumeService:
    """负责简历格式校验、文本提取、元数据和文本文件持久化。"""

    MAX_BYTES = 10 * 1024 * 1024
    MAX_CHARACTERS = 120_000
    MODEL_CHARACTERS = 60_000
    ALLOWED_SUFFIXES = {".pdf", ".docx"}
    _ID_PATTERN = re.compile(r"^resume-[a-f0-9]{32}$")
    _OCR_ENGINE = None

    def __init__(self, storage_path: str | Path | None = None, state_store=None) -> None:
        settings = get_settings()
        self.storage_path = Path(storage_path or settings.resume_storage_path)
        self.state_store = state_store or LocalStateStore()

    def save(self, filename: str | None, data: bytes) -> dict:
        """解析一份 PDF/DOCX 简历并原子保存其纯文本。"""
        safe_name = Path(filename or "resume").name
        suffix = Path(safe_name).suffix.lower()
        if suffix not in self.ALLOWED_SUFFIXES:
            raise ApplicationError(
                "UNSUPPORTED_RESUME_FORMAT",
                "仅支持 PDF 或 DOCX 格式的简历；旧版 DOC 请先另存为 DOCX。",
                status_code=415,
            )
        if not data:
            raise ApplicationError("EMPTY_RESUME", "简历文件为空。")
        if len(data) > self.MAX_BYTES:
            raise ApplicationError("RESUME_TOO_LARGE", "简历文件不能超过 10 MB。", status_code=413)

        if suffix == ".pdf":
            text, parse_method, quality_score = self._parse_pdf(data)
        else:
            text = self._normalize(self._parse_docx(data))
            quality_score, readable = self._text_quality(text)
            parse_method = "python-docx"
            if not readable:
                raise ApplicationError(
                    "RESUME_TEXT_UNREADABLE",
                    "Word 简历中的文本无法可靠识别，请检查文件内容后重新上传。",
                )
        if len(text) < 20:
            raise ApplicationError(
                "RESUME_TEXT_NOT_FOUND",
                "没有从简历中读取到足够文本；扫描版 PDF 请先进行 OCR。",
            )
        if len(text) > self.MAX_CHARACTERS:
            raise ApplicationError(
                "RESUME_TEXT_TOO_LONG",
                "简历解析后的文本过长，请精简到 12 万字符以内后重新上传。",
                status_code=413,
            )
        resume_id = "resume-" + uuid4().hex
        created_at = datetime.now().isoformat()
        record = {
            "id": resume_id,
            "file_name": safe_name,
            "format": suffix.removeprefix("."),
            "character_count": len(text),
            "parse_method": parse_method,
            "quality_score": quality_score,
            "created_at": created_at,
        }
        self.storage_path.mkdir(parents=True, exist_ok=True)
        target = self.storage_path / f"{resume_id}.txt"
        temporary = target.with_suffix(".txt.tmp")
        try:
            temporary.write_text(text, encoding="utf-8")
            os.replace(temporary, target)
        except OSError as exc:
            temporary.unlink(missing_ok=True)
            target.unlink(missing_ok=True)
            raise ApplicationError(
                "RESUME_SAVE_FAILED", "简历文本保存失败，请检查本地目录权限。", status_code=500
            ) from exc
        try:
            self.state_store.put("resumes", resume_id, record)
        except Exception:
            target.unlink(missing_ok=True)
            raise
        return record

    def get(self, resume_id: str) -> dict:
        """返回已保存的简历元数据，不暴露本地文件路径。"""
        if not self._ID_PATTERN.fullmatch(resume_id):
            raise ApplicationError("RESUME_NOT_FOUND", "简历不存在，请重新上传。", status_code=404)
        record = self.state_store.get("resumes", resume_id)
        if record is None or not (self.storage_path / f"{resume_id}.txt").is_file():
            raise ApplicationError("RESUME_NOT_FOUND", "简历不存在，请重新上传。", status_code=404)
        return record

    def model_context(self, resume_id: str) -> dict:
        """返回供面试模型分析的简历文本，并限制提示词体积。"""
        record = self.get(resume_id)
        try:
            text = (self.storage_path / f"{resume_id}.txt").read_text(encoding="utf-8")
        except OSError as exc:
            raise ApplicationError("RESUME_READ_FAILED", "简历文本读取失败。", status_code=500) from exc
        return {
            "file_name": record["file_name"],
            "text": text[: self.MODEL_CHARACTERS],
            "truncated": len(text) > self.MODEL_CHARACTERS,
        }

    @classmethod
    def _parse_pdf(cls, data: bytes) -> tuple[str, str, float]:
        """选择质量最高的全文提取结果，必要时对整页图像执行本地 OCR。"""
        candidates: list[tuple[float, str, str]] = []
        failures = []
        for method, parser in (
            ("pypdf", cls._parse_pdf_pypdf),
            ("pymupdf", cls._parse_pdf_pymupdf),
        ):
            try:
                text = cls._normalize(parser(data))
                score, readable = cls._text_quality(text)
                if text:
                    candidates.append((score, method, text))
                if not readable:
                    continue
            except ApplicationError as exc:
                if exc.code == "ENCRYPTED_RESUME":
                    raise
                failures.append(exc)
            except Exception as exc:
                failures.append(exc)

        readable_candidates = [candidate for candidate in candidates if cls._text_quality(candidate[2])[1]]
        if readable_candidates:
            score, method, text = max(
                readable_candidates, key=lambda candidate: (candidate[0], len(candidate[2]))
            )
            return text, method, score

        try:
            text = cls._normalize(cls._parse_pdf_ocr(data))
            score, readable = cls._text_quality(text)
            if text:
                candidates.append((score, "rapidocr", text))
            if readable:
                return text, "rapidocr", score
        except ApplicationError as exc:
            if exc.code == "ENCRYPTED_RESUME":
                raise
            failures.append(exc)
        except Exception as exc:
            failures.append(exc)

        if not candidates and failures:
            raise ApplicationError(
                "INVALID_RESUME", "PDF 简历无法解析，请检查文件是否损坏。"
            ) from failures[-1]
        raise ApplicationError(
            "RESUME_TEXT_UNREADABLE",
            "PDF 的字体编码异常，未能可靠识别完整文字。请改用 DOCX 或可复制文字的 PDF。",
        )

    @staticmethod
    def _parse_pdf_pypdf(data: bytes) -> str:
        try:
            reader = PdfReader(BytesIO(data), strict=False)
            if reader.is_encrypted and reader.decrypt("") == 0:
                raise ApplicationError("ENCRYPTED_RESUME", "无法解析加密的 PDF 简历。")
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except ApplicationError:
            raise
        except Exception as exc:
            raise ApplicationError("INVALID_RESUME", "PDF 简历无法解析，请检查文件是否损坏。") from exc

    @staticmethod
    def _parse_pdf_pymupdf(data: bytes) -> str:
        """使用另一套 PDF 字体映射实现提取全部页面，修复部分中文 CMap 问题。"""
        try:
            with pymupdf.open(stream=data, filetype="pdf") as document:
                if document.needs_pass:
                    raise ApplicationError("ENCRYPTED_RESUME", "无法解析加密的 PDF 简历。")
                return "\n".join(page.get_text("text", sort=True) for page in document)
        except ApplicationError:
            raise
        except Exception as exc:
            raise ApplicationError("INVALID_RESUME", "PDF 简历无法解析，请检查文件是否损坏。") from exc

    @classmethod
    def _parse_pdf_ocr(cls, data: bytes) -> str:
        """将每一页渲染为图像后离线 OCR，避免依赖 PDF 内部错误的字符映射。"""
        try:
            if cls._OCR_ENGINE is None:
                from rapidocr import RapidOCR

                cls._OCR_ENGINE = RapidOCR()
            pages = []
            with pymupdf.open(stream=data, filetype="pdf") as document:
                if document.needs_pass:
                    raise ApplicationError("ENCRYPTED_RESUME", "无法解析加密的 PDF 简历。")
                for page in document:
                    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), colorspace=pymupdf.csRGB, alpha=False)
                    image = np.frombuffer(pixmap.samples, dtype=np.uint8).reshape(
                        pixmap.height, pixmap.width, pixmap.n
                    )
                    result = cls._OCR_ENGINE(image)
                    lines = tuple(result.txts or ())
                    pages.append("\n".join(lines))
            return "\n".join(pages)
        except ApplicationError:
            raise
        except Exception as exc:
            raise ApplicationError("RESUME_OCR_FAILED", "PDF 简历 OCR 识别失败。") from exc

    @staticmethod
    def _parse_docx(data: bytes) -> str:
        try:
            document = Document(BytesIO(data))
            blocks = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                blocks.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
            return "\n".join(blocks)
        except Exception as exc:
            raise ApplicationError("INVALID_RESUME", "Word 简历无法解析，请检查文件是否损坏。") from exc

    @staticmethod
    def _normalize(text: str) -> str:
        lines = []
        for line in text.replace("\x00", "").splitlines():
            normalized = " ".join(line.split())
            if normalized:
                lines.append(normalized)
        return "\n".join(lines)

    @staticmethod
    def _text_quality(text: str) -> tuple[float, bool]:
        """识别乱码而不依赖公司名或简历模板，避免把错误 Unicode 当成有效全文。"""
        if len(text) < 20:
            return 0.0, False

        letters = 0
        digits = 0
        invalid = 0
        scripts: dict[str, int] = {}
        for character in text:
            codepoint = ord(character)
            category = unicodedata.category(character)
            if character == "\ufffd" or 0xE000 <= codepoint <= 0xF8FF:
                invalid += 1
            elif category.startswith("C") and character not in "\n\r\t":
                invalid += 1
            if character.isdigit():
                digits += 1
            if not character.isalpha():
                continue
            letters += 1
            name = unicodedata.name(character, "UNKNOWN")
            if "CJK" in name or "IDEOGRAPH" in name:
                script = "HAN"
            else:
                script = name.split(" ", 1)[0]
            scripts[script] = scripts.get(script, 0) + 1

        meaningful = letters + digits
        if meaningful < 12:
            return 0.1, False
        active_threshold = max(3, round(letters * 0.02))
        active_scripts = [count for count in scripts.values() if count >= active_threshold]
        invalid_ratio = invalid / max(len(text), 1)
        script_penalty = max(0, len(active_scripts) - 3) * 0.18
        symbol_penalty = 0.35 if meaningful / len(text) < 0.25 else 0.0
        quality = max(0.0, min(1.0, 1.0 - invalid_ratio * 8 - script_penalty - symbol_penalty))
        readable = quality >= 0.65 and invalid_ratio < 0.02 and len(active_scripts) <= 4
        return round(quality, 3), readable
