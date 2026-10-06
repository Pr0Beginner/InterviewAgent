"""本地业务文件、Skill 服务和邮件去重的回归测试。"""

import asyncio
import tempfile
import unittest
from contextlib import contextmanager
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from backend.app.agent.workflows import InterviewWorkflows
from backend.app.core.exceptions import ApplicationError
from backend.app.schemas.business import EmailEvent, JobSearchRequest, MockAnswerRequest, MockInterviewRequest
from backend.app.services.email_sync import EmailSyncService
from backend.app.services.jobs import JobService
from backend.app.services.mock_interviews import MockInterviewService
from backend.app.services.resumes import ResumeService
from backend.app.storage.local import LocalStateStore, MarkdownInterviewStore


class FakeModel:
    def __init__(self, results):
        self.results = iter(results)
        self.calls = 0
        self.payloads = []

    def ensure_configured(self):
        pass

    async def structured(self, prompt, data, schema):
        self.calls += 1
        self.payloads.append(data)
        return schema.model_validate(next(self.results))


class StoreTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        path = Path(self.temp.name)
        self.store = MarkdownInterviewStore(path / "interviews.md")
        self.state = LocalStateStore(path / "runtime.json")
        self.graphs = InterviewWorkflows(self.store, path / "graph.sqlite", self.state)

    def create(self):
        return self.graphs.mutate({
            "company_name": "测试公司", "position_name": "Java 后端",
            "base_location": "杭州", "current_status": "一面",
        }, "create-1", create=True)


class WorkflowsTest(StoreTest):
    def test_create_update_query_and_replay_use_local_files(self):
        created = self.create()
        args = {"interview_id": created["id"], "expected_status": "一面", "target_status": "二面", "interview_time": "2026-10-03T14:00:00"}
        changed = self.graphs.mutate(args, "update-1")
        replay = self.graphs.mutate(args, "update-1")
        self.assertEqual(changed["current_status"], "二面")
        self.assertEqual(replay, changed)
        self.assertEqual(self.graphs.query({"company_name": "测试公司"})["source"], "local_markdown")
        self.assertTrue((Path(self.temp.name) / "interviews.md").is_file())

    def test_stale_status_cannot_overwrite_recent_change(self):
        created = self.create()
        with self.assertRaisesRegex(ApplicationError, "重新查询"):
            self.graphs.mutate({"interview_id": created["id"], "expected_status": "笔试中", "target_status": "三面"}, "stale")
        self.assertEqual(self.store.get(created["id"]).current_status, "一面")

    def test_agent_can_update_any_record_field(self):
        created = self.create()
        changed = self.graphs.mutate({
            "interview_id": created["id"], "expected_status": "一面",
            "position_name": "AI应用研发", "base_location": "深圳",
            "job_url": "https://example.com/jobs/ai-application",
        }, "update-fields")
        self.assertEqual(changed["position_name"], "AI应用研发")
        self.assertEqual(changed["base_location"], "深圳")
        self.assertEqual(self.store.get(created["id"]).job_url, "https://example.com/jobs/ai-application")


class SkillServicesTest(StoreTest):
    def test_pdf_resume_text_is_extracted(self):
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({
                NameObject("/F1"): writer._add_object(font),
            }),
        })
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 72 720 Td (Java Spring Boot resume project) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(content)
        buffer = BytesIO()
        writer.write(buffer)
        service = ResumeService(Path(self.temp.name) / "resumes", self.state)

        saved = service.save("resume.pdf", buffer.getvalue())

        self.assertIn("Spring Boot", service.model_context(saved["id"])["text"])

    def test_resume_is_parsed_locally_and_injected_into_interview_prompt(self):
        document = Document()
        document.add_heading("张三 - Java 后端", level=1)
        document.add_paragraph("技术栈：Java、Spring Boot、MySQL、Redis")
        document.add_paragraph("项目：订单系统。负责缓存一致性设计，将 P99 延迟降低 35%。")
        buffer = BytesIO()
        document.save(buffer)
        resume_service = ResumeService(Path(self.temp.name) / "resumes", self.state)
        saved = resume_service.save("张三简历.docx", buffer.getvalue())
        provider = FakeModel([{"question": "订单系统的缓存一致性方案为什么这样设计？"}])
        service = MockInterviewService(
            self.state,
            provider,
            resume_service=resume_service,
        )

        created = asyncio.run(service.create(MockInterviewRequest(
            position_name="Java 后端",
            interview_round="一面",
            resume_id=saved["id"],
        )))

        self.assertEqual(created["resume"]["file_name"], "张三简历.docx")
        self.assertIn("Spring Boot", provider.payloads[0]["context"]["resume"]["text"])
        stored_context = service.get(created["session_id"])["context"]
        self.assertEqual(stored_context["resume_id"], saved["id"])
        self.assertNotIn("resume", stored_context)
        self.assertTrue((Path(self.temp.name) / "resumes" / f"{saved['id']}.txt").is_file())

    def test_mock_interview_lifecycle_persists_in_json(self):
        provider = FakeModel([
            {"question": "解释 Java 的 volatile 语义。"},
            {"evaluation": "说明了可见性，缺少有序性。", "score": 65, "question": "volatile 能保证 i++ 原子性吗？", "is_follow_up": True},
            {"overall_score": 65, "summary": "基础概念尚需补全。", "strengths": ["知道可见性"], "weaknesses": ["遗漏有序性"], "suggestions": ["复习 happens-before"]},
        ])
        service = MockInterviewService(self.state, provider)
        created = asyncio.run(service.create(MockInterviewRequest(company_name="网易", position_name="Java 后端", interview_round="一面")))
        request = MockAnswerRequest(question_id=created["question_id"], answer="保证线程间可见性。")
        first = asyncio.run(service.answer(created["session_id"], request))
        self.assertEqual(asyncio.run(service.answer(created["session_id"], request)), first)
        self.assertIn("answer", MockInterviewService(self.state, provider).get(created["session_id"])["turns"][0])
        report = asyncio.run(service.finish(created["session_id"]))
        self.assertEqual(asyncio.run(service.finish(created["session_id"])), report)

    def test_job_pagination_reuses_local_snapshot(self):
        class PreferenceMemory:
            def read_summary(self):
                return "偏好杭州的稳定后端业务。"

        class Source:
            calls = 0
            async def search(self, keywords, limit, cities=None, work_experience="应届生"):
                self.calls += 1
                self.cities = cities
                self.work_experience = work_experience
                return [{"job_url": f"https://www.zhipin.com/job_detail/{index}.html", "job_description": "Java 后端，杭州，MySQL", "salary": "20-30K"} for index in range(3)]
        source = Source()
        provider = FakeModel([{"items": [
            {"source_index": index, "company_name": "测试公司", "position_name": "Java 后端", "base_location": "杭州", "match_score": 90 - index, "match_reasons": ["Java 匹配"], "risk_points": ["经验待核对"], "jd_summary": "Java 后端"}
            for index in range(3)
        ]}])
        service = JobService(self.state, provider, source, PreferenceMemory())
        request = JobSearchRequest(cities=["杭州"], keywords=["Java 后端"], page_size=2)
        page1 = asyncio.run(service.recommend(request))
        request.page = 2
        page2 = asyncio.run(service.recommend(request))
        self.assertEqual(len(page2["items"]), 1)
        self.assertEqual(source.calls, 1)
        self.assertEqual(source.cities, ["杭州"])
        self.assertEqual(source.work_experience, "应届生")
        self.assertEqual(provider.calls, 1)
        self.assertEqual(len(provider.payloads[0]["sources"]), 3)
        self.assertEqual(page1["items"][0]["salary"], "20-30K")
        self.assertEqual(provider.payloads[0]["profile"]["preference_memory"], "偏好杭州的稳定后端业务。")
        self.assertEqual(service.get(page1["items"][0]["id"])["job_url"], "https://www.zhipin.com/job_detail/0.html")


class MailServiceTest(StoreTest):
    def test_missing_local_record_becomes_approvable_candidate(self):
        """高置信邮件找不到本地记录时，必须等待用户确认后再创建。"""
        class Mailbox:
            marks = 0
            def ensure_configured(self): pass
            @contextmanager
            def connect(self): yield self
            def unread(self, client, limit):
                return [{"uid": "88", "receipt_id": "b" * 64}]
            def fetch(self, client, uid):
                return {
                    "subject": "卓望公司线上笔试通知",
                    "body": "卓望公司邀请你参加全栈开发工程师线上笔试，工作地点广州。",
                }
            def mark_read(self, client, uid):
                self.marks += 1

        provider = FakeModel([EmailEvent(
            relevant=True,
            confidence=1,
            company_name="卓望公司",
            position_name="全栈开发工程师",
            base_location="广州",
            status="笔试中",
            evidence="卓望公司邀请你参加全栈开发工程师线上笔试",
        ).model_dump()])
        mailbox = Mailbox()
        service = EmailSyncService(self.state, self.store, mailbox, provider, self.graphs)
        task = service.create_task(1)
        service.run(task["task_id"], 1)

        candidates = service.get_task(task["task_id"])["result"]["creation_candidates"]
        self.assertEqual(len(candidates), 1)
        self.assertEqual(candidates[0]["company_name"], "卓望公司")
        self.assertEqual(self.store.all(), [])
        self.assertEqual(mailbox.marks, 0)

        created = service.create_candidates([candidates[0]["id"]])
        self.assertEqual(len(created["created"]), 1)
        self.assertEqual(self.store.all()[0].company_name, "卓望公司")
        repeated = service.create_candidates([candidates[0]["id"]])
        self.assertEqual(len(repeated["already_created"]), 1)
        self.assertEqual(len(self.store.all()), 1)

    def test_processed_mail_retry_does_not_repeat_business_write(self):
        created = self.create()

        class Mailbox:
            marks = 0
            fail_mark = True
            def ensure_configured(self): pass
            @contextmanager
            def connect(self): yield self
            def unread(self, client, limit): return [{"uid": "42", "receipt_id": "a" * 64}]
            def fetch(self, client, uid): return {"subject": "测试公司二面邀请", "body": "测试公司 Java 后端二面邀请，时间 2026-10-08 14:00。"}
            def mark_read(self, client, uid):
                self.marks += 1
                if self.fail_mark:
                    self.fail_mark = False
                    raise RuntimeError("mark failed")

        provider = FakeModel([EmailEvent(relevant=True, confidence=1, company_name="测试公司", position_name="Java 后端", status="二面", interview_time="2026-10-08T14:00:00", evidence="测试公司 Java 后端二面邀请").model_dump()])
        mailbox = Mailbox()
        service = EmailSyncService(self.state, self.store, mailbox, provider, self.graphs)
        first = service.create_task(1)
        service.run(first["task_id"], 1)
        second = service.create_task(1)
        service.run(second["task_id"], 1)
        self.assertEqual(self.store.get(created["id"]).current_status, "二面")
        self.assertEqual(provider.calls, 1)
        self.assertEqual(service.get_task(second["task_id"])["result"]["already_processed"], 1)


if __name__ == "__main__":
    unittest.main()
